"""Dagelijkse rapporten in de map rapporten/.

Na elke ronde:
- rapporten/AAAA-MM-DD.md   een notitie van die dag (blijft bewaard);
- rapporten/trend-agent.xlsx een Excel-bestand dat elke dag opnieuw wordt
  opgebouwd uit de nep-portefeuille, met vijf tabbladen:
    Overzicht   waarde per dag, dagresultaat en totaalresultaat
    Posities    wat de agent nu in bezit heeft
    Trades      elke aan- en verkoop met reden en resultaat
    Trends      de sterkste trends van de laatste ronde
    Uitleg      wat alle kolommen betekenen

Het Excel-bestand wordt altijd volledig uit portfolio.json opgebouwd, zodat
het nooit uit de pas kan lopen met wat de agent werkelijk heeft gedaan.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .portfolio import Portfolio
from .universe import BY_TICKER

HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
HEADER_FONT = Font(bold=True, color="FFFFFF")
EUR = '€ #,##0.00;[Red]-€ #,##0.00'
PCT = '0.00%;[Red]-0.00%'


def _sheet(wb: Workbook, title: str, headers: list[str], widths: list[int]):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for col, (h, w) in enumerate(zip(headers, widths), start=1):
        cell = ws.cell(row=1, column=col)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(vertical="center")
        ws.column_dimensions[get_column_letter(col)].width = w
    ws.freeze_panes = "A2"
    return ws


def write_excel(pf: Portfolio, ranking, path: Path) -> None:
    wb = Workbook()
    wb.remove(wb.active)

    # --- Overzicht: waarde per dag, met formules zodat je ze zelf kunt narekenen
    ws = _sheet(wb, "Overzicht", ["Datum", "Waarde", "Dagresultaat", "Dag %", "Totaal %"], [14, 16, 16, 10, 10])
    for i, e in enumerate(pf.equity_history, start=2):
        ws.cell(row=i, column=1, value=e["day"])
        ws.cell(row=i, column=2, value=e["equity"]).number_format = EUR
        if i > 2:
            ws.cell(row=i, column=3, value=f"=B{i}-B{i - 1}").number_format = EUR
            ws.cell(row=i, column=4, value=f"=B{i}/B{i - 1}-1").number_format = PCT
        ws.cell(row=i, column=5, value=f"=B{i}/$B$2-1").number_format = PCT

    # --- Posities
    ws = _sheet(wb, "Posities", ["Ticker", "Naam", "Regio", "Aantal", "Instapprijs", "Stop-loss",
                                 "Valuta", "Sinds", "Max verlies (€)"], [11, 22, 7, 8, 12, 12, 7, 12, 14])
    for p in pf.positions.values():
        inst = BY_TICKER.get(p.ticker)
        ws.append([p.ticker, inst.name if inst else "", inst.region if inst else "", p.qty,
                   round(p.entry_price, 2), round(p.stop, 2), p.currency, p.opened, p.risk_eur])
        ws.cell(row=ws.max_row, column=9).number_format = EUR
    if not pf.positions:
        ws.append(["(geen open posities)"])

    # --- Trades
    ws = _sheet(wb, "Trades", ["Datum", "Actie", "Ticker", "Aantal", "Prijs", "Valuta-koers",
                               "Bedrag (€)", "Kosten (€)", "Resultaat (€)", "Reden"],
                [12, 8, 11, 8, 11, 12, 13, 11, 13, 70])
    for t in pf.trades:
        ws.append([t.day, "KOOP" if t.side == "BUY" else "VERKOOP", t.ticker, t.qty, t.price,
                   round(t.fx, 4), t.value_eur, t.costs_eur, t.pnl_eur, t.reason])
        for col in (7, 8, 9):
            ws.cell(row=ws.max_row, column=col).number_format = EUR
        ws.cell(row=ws.max_row, column=5).number_format = '0.00'
    if not pf.trades:
        ws.append(["(nog geen trades)"])

    # --- Trends van de laatste ronde
    ws = _sheet(wb, "Trends", ["#", "Ticker", "Naam", "Regio", "Score", "3 maanden", "1 maand",
                               "Volume t.o.v. normaal", "Slotkoers"], [5, 11, 22, 7, 9, 11, 10, 20, 11])
    for i, a in enumerate(ranking, start=1):
        inst = BY_TICKER.get(a.ticker)
        ws.append([i, a.ticker, inst.name if inst else "", inst.region if inst else "",
                   round(a.score, 3), a.ret63, a.ret21, round(a.rel_volume, 2), round(a.close, 2)])
        ws.cell(row=ws.max_row, column=6).number_format = PCT
        ws.cell(row=ws.max_row, column=7).number_format = PCT
    if not ranking:
        ws.append(["", "(geen kandidaten die aan alle regels voldoen)"])

    # --- Uitleg
    ws = _sheet(wb, "Uitleg", ["Begrip", "Betekenis"], [24, 100])
    for row in [
        ("Nepgeld", "Dit is een oefenportefeuille. Er is geen echt geld gebruikt. Geen beleggingsadvies."),
        ("Waarde", "Kas plus de waarde van alle posities, omgerekend naar euro."),
        ("Stop-loss", "Koers waarop de agent automatisch verkoopt om verlies te beperken. Schuift mee omhoog als het goed gaat."),
        ("Max verlies (€)", "Wat de positie kost als de beginstop wordt geraakt, inclusief kosten."),
        ("Score", "Trendsterkte: rendement over 1 en 3 maanden gedeeld door beweeglijkheid. Hoger = sterker en rustiger."),
        ("Volume t.o.v. normaal", "Handel van de laatste 20 dagen gedeeld door de laatste 60. Boven 1 = meer belangstelling."),
        ("Resultaat (€)", "Winst of verlies van een afgesloten trade, na alle kosten."),
    ]:
        ws.append(list(row))
        ws.cell(row=ws.max_row, column=2).alignment = Alignment(wrap_text=True)

    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_suffix(".tmp.xlsx")
    wb.save(tmp)
    tmp.replace(path)


def write_daily_note(day: str, markdown: str, folder: Path) -> Path:
    folder.mkdir(parents=True, exist_ok=True)
    path = folder / f"{day}.md"
    path.write_text(markdown, encoding="utf-8")
    return path


def write_investor_excel(book, out: dict, path: Path) -> None:
    """Excel-rapport van de investeerder: waarde per dag, verdeling en trades."""
    wb = Workbook()
    wb.remove(wb.active)
    ws = _sheet(wb, "Overzicht", ["Datum", "Waarde", "Dagresultaat", "Dag %", "Totaal %"], [14, 16, 16, 10, 10])
    for i, e in enumerate(book.equity_history, start=2):
        ws.cell(row=i, column=1, value=e["day"])
        ws.cell(row=i, column=2, value=e["equity"]).number_format = EUR
        if i > 2:
            ws.cell(row=i, column=3, value=f"=B{i}-B{i - 1}").number_format = EUR
            ws.cell(row=i, column=4, value=f"=B{i}/B{i - 1}-1").number_format = PCT
        ws.cell(row=i, column=5, value=f"=B{i}/$B$2-1").number_format = PCT

    ws = _sheet(wb, "Verdeling", ["Ticker", "Naam", "Soort", "Aandeel"], [11, 24, 9, 10])
    for t, w in sorted(out.get("weights", {}).items(), key=lambda kv: -kv[1]):
        inst = BY_TICKER.get(t)
        ws.append([t, inst.name if inst else "", inst.kind if inst else "", w])
        ws.cell(row=ws.max_row, column=4).number_format = PCT
    if out.get("equity_eur"):
        ws.append(["KAS", "", "", out.get("cash_eur", 0) / out["equity_eur"]])
        ws.cell(row=ws.max_row, column=4).number_format = PCT

    ws = _sheet(wb, "Trades", ["Datum", "Actie", "Ticker", "Aantal", "Prijs (€)", "Bedrag (€)",
                               "Kosten (€)", "Resultaat (€)", "Reden"], [12, 8, 10, 11, 11, 12, 10, 13, 40])
    for t in book.trades:
        ws.append([t.day, "KOOP" if t.side == "BUY" else "VERKOOP", t.ticker, t.qty, t.price,
                   t.value_eur, t.costs_eur, t.pnl_eur, t.reason])
        for col in (6, 7, 8):
            ws.cell(row=ws.max_row, column=col).number_format = EUR

    ws = _sheet(wb, "Uitleg", ["Begrip", "Betekenis"], [20, 100])
    for row in [
        ("Nepgeld", "Oefenportefeuille, geen echt geld. Geen beleggingsadvies."),
        ("Kern (60%)", "Wereldwijde ETF (MSCI World). Volgt de hele markt."),
        ("Satelliet (30%)", "De 5 sterkste aandelen/ETF's over 12 maanden die in een stijgende trend zitten."),
        ("Crypto (10%)", "Bitcoin en Ethereum, alleen boven hun 200-daags gemiddelde. Anders kas."),
        ("Herverdelen", "Eén keer per maand, op de eerste handelsdag."),
        ("Aantal", "Gebroken aantallen zijn toegestaan; niet elke echte broker ondersteunt dat."),
    ]:
        ws.append(list(row))
        ws.cell(row=ws.max_row, column=2).alignment = Alignment(wrap_text=True)
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_suffix(".tmp.xlsx")
    wb.save(tmp)
    tmp.replace(path)
