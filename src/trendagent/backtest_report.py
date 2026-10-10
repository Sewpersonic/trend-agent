"""Schrijft de uitkomst van een backtest naar backtest/resultaten/:
- RAPPORT.md       leesbaar verslag (GitHub toont dit netjes)
- backtest.xlsx    alle cijfers, per jaar, per 4 weken, waardeverloop en trades
- waardeverloop.png grafiek van agent tegenover de ETF's (als matplotlib er is)
"""

from __future__ import annotations

from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .backtest import (WINDOW_DAYS, BacktestResult, per_year, summary, trade_stats,
                       window_stats, windows)

NAMES = {"AGENT": "Trend-agent", "IWDA.AS": "MSCI World-ETF (IWDA)", "SPY": "S&P 500-ETF (SPY, in euro)"}


def _pct(x: float) -> str:
    return f"{x:+.1%}".replace(".", ",")


def _eur(x: float) -> str:
    return f"€{x:,.0f}".replace(",", ".")


def _all_series(r: BacktestResult) -> dict[str, list]:
    return {"AGENT": r.equity, **r.benchmarks}


def verdict(r: BacktestResult) -> list[str]:
    """Een paar eerlijke zinnen bovenaan het rapport."""
    a = summary(r.equity)
    w = window_stats(windows(r.equity))
    lines = []
    bench = r.benchmarks.get("IWDA.AS") or next(iter(r.benchmarks.values()), None)
    if bench:
        b = summary(bench)
        if a["total"] > b["total"]:
            lines.append(f"Over de hele periode deed de agent het beter dan gewoon de wereld-ETF vasthouden "
                         f"({_pct(a['total'])} tegen {_pct(b['total'])}).")
        else:
            lines.append(f"Over de hele periode deed de agent het slechter dan gewoon de wereld-ETF vasthouden "
                         f"({_pct(a['total'])} tegen {_pct(b['total'])}).")
        lines.append(f"Zijn grootste daling was {_pct(a['max_dd'])}, tegen {_pct(b['max_dd'])} voor de ETF.")
    if w:
        lines.append(f"In {w['positive_share']:.0%} van alle periodes van 4 weken stond de agent aan het eind hoger dan "
                     f"aan het begin. Het slechtste blok van 4 weken was {_pct(w['worst'])}, het beste {_pct(w['best'])}.")
    return lines


def write_markdown(r: BacktestResult, path: Path, years: int) -> None:
    t = trade_stats(r.trades)
    rows = []
    for key, series in _all_series(r).items():
        s = summary(series)
        if s:
            rows.append(f"| {NAMES.get(key, key)} | {_pct(s['total'])} | {_pct(s['cagr'])} | "
                        f"{_pct(s['max_dd'])} | {s["sharpe"]:.2f} | {_eur(series[-1][1])} |")
    yr = {k: per_year(v) for k, v in _all_series(r).items()}
    years_list = sorted(yr["AGENT"])
    yr_rows = [f"| {y} | " + " | ".join(_pct(yr[k].get(y, 0.0)) if y in yr[k] else "–" for k in yr) + " |"
               for y in years_list]
    win = {k: window_stats(windows(v)) for k, v in _all_series(r).items()}

    lines = [
        "# Backtest trend-agent",
        "",
        f"Periode **{r.start} t/m {r.end}** (ongeveer {years} jaar), startkapitaal €10.000 nepgeld, "
        f"{r.tickers_used} instrumenten met koersdata. Zelfde regels, kosten en slippage als bij het dagelijks oefenen.",
        "",
        "> Dit is geen voorspelling. Het laat zien hoe deze regels het in het verleden hadden gedaan.",
        "",
        "## In het kort",
        "",
        *[f"- {l}" for l in verdict(r)],
        "",
        "## Totaal",
        "",
        "| | Totaal | Per jaar | Grootste daling | Sharpe | Eindwaarde |",
        "| --- | --- | --- | --- | --- | --- |",
        *rows,
        "",
        "_Sharpe: rendement gedeeld door beweeglijkheid. Boven 1 is goed, onder 0,5 matig._",
        "",
        "## Per kalenderjaar",
        "",
        "| Jaar | " + " | ".join(NAMES.get(k, k) for k in yr) + " |",
        "| --- | " + " | ".join("---" for _ in yr) + " |",
        *yr_rows,
        "",
        f"## Periodes van 4 weken ({WINDOW_DAYS} handelsdagen)",
        "",
        "Dit beantwoordt de vraag \"wat kan er in een paar weken gebeuren?\": elk mogelijk blok van 4 weken in de periode.",
        "",
        "| | Hoger geëindigd | Normaal (mediaan) | Slecht (1 op 20) | Slechtste | Beste |",
        "| --- | --- | --- | --- | --- | --- |",
        *[f"| {NAMES.get(k, k)} | {w['positive_share']:.0%} | {_pct(w['median'])} | {_pct(w['p05'])} | "
          f"{_pct(w['worst'])} | {_pct(w['best'])} |" for k, w in win.items() if w],
        "",
        "## Trades",
        "",
        f"- Aankopen: {t['buys']}, afgesloten: {t['closed']}, waarvan winst: {t['win_rate']:.0%}",
        f"- Gemiddelde winst: {_eur(t['avg_win'])}, gemiddeld verlies: {_eur(t['avg_loss'])}, "
        f"profit factor: {str(round(t['profit_factor'], 2)).replace('.', ',')}",
        f"- Gemiddeld vastgehouden: {t['avg_hold_days']:.0f} dagen; totale kosten: {_eur(t['costs'])}",
        f"- Kill switch aangegaan: {len(r.halts)} keer"
        + (" (" + "; ".join(f"{d}: {why}" for d, why in r.halts[:6]) + (" …" if len(r.halts) > 6 else "") + ")" if r.halts else ""),
        "",
        "## Belangrijke kanttekeningen",
        "",
        "- **Survivorship bias.** De lijst instrumenten is in 2026 gekozen. Aandelen die de afgelopen jaren "
        "groot zijn geworden (Nvidia, Palantir, Rheinmetall) staan er daardoor in, en verliezers die uit de "
        "indexen zijn verdwenen niet. De echte uitkomst was waarschijnlijk minder goed.",
        "- **Klein risico per trade.** Met 0,25% risico per trade is de agent vaak maar voor een deel belegd. "
        "Daardoor zijn zowel winst als verlies kleiner dan bij de ETF.",
        "- **Gratis koersdata** van Yahoo Finance kan fouten of gaten bevatten.",
        "- **Na de kill switch** doet de backtest alsof jij na 5 handelsdagen 'resume' doet.",
        "- Een goed resultaat in het verleden is geen toestemming om met echt geld te handelen.",
        "",
        "![Waardeverloop](waardeverloop.png)",
        "",
    ]
    path.write_text("\n".join(lines), encoding="utf-8")


def write_excel(r: BacktestResult, path: Path) -> None:
    head_fill, head_font = PatternFill("solid", fgColor="1F3A5F"), Font(bold=True, color="FFFFFF")
    pct, eur = "0.0%;[Red]-0.0%", "€ #,##0;[Red]-€ #,##0"
    wb = Workbook()
    wb.remove(wb.active)

    def sheet(title, headers, widths):
        ws = wb.create_sheet(title)
        ws.append(headers)
        for i, w in enumerate(widths, 1):
            ws.cell(1, i).fill, ws.cell(1, i).font = head_fill, head_font
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "A2"
        return ws

    series = _all_series(r)
    ws = sheet("Samenvatting", ["", "Totaal", "Per jaar", "Grootste daling", "Sharpe", "Eindwaarde"],
               [28, 11, 11, 16, 9, 13])
    for k, v in series.items():
        s = summary(v)
        if s:
            ws.append([NAMES.get(k, k), s["total"], s["cagr"], s["max_dd"], round(s["sharpe"], 2), v[-1][1]])
            for c in (2, 3, 4):
                ws.cell(ws.max_row, c).number_format = pct
            ws.cell(ws.max_row, 6).number_format = eur

    yr = {k: per_year(v) for k, v in series.items()}
    ws = sheet("Per jaar", ["Jaar"] + [NAMES.get(k, k) for k in yr], [8] + [26] * len(yr))
    for y in sorted(yr["AGENT"]):
        ws.append([y] + [yr[k].get(y) for k in yr])
        for c in range(2, len(yr) + 2):
            ws.cell(ws.max_row, c).number_format = pct

    ws = sheet("Per 4 weken", ["", "Aantal blokken", "Hoger geëindigd", "Mediaan", "1 op 20 slecht",
                               "Slechtste", "Beste"], [28, 14, 15, 10, 14, 11, 10])
    for k, v in series.items():
        w = window_stats(windows(v))
        if w:
            ws.append([NAMES.get(k, k), w["count"], w["positive_share"], w["median"], w["p05"], w["worst"], w["best"]])
            for c in range(3, 8):
                ws.cell(ws.max_row, c).number_format = pct

    keys = list(series)
    lookup = {k: dict(v) for k, v in series.items()}
    ws = sheet("Waardeverloop", ["Datum"] + [NAMES.get(k, k) for k in keys], [12] + [24] * len(keys))
    for d, _ in r.equity:
        ws.append([d.isoformat()] + [lookup[k].get(d) for k in keys])
        for c in range(2, len(keys) + 2):
            ws.cell(ws.max_row, c).number_format = eur

    ws = sheet("Trades", ["Datum", "Actie", "Ticker", "Aantal", "Prijs", "Bedrag (€)", "Kosten (€)",
                          "Resultaat (€)", "Reden"], [12, 8, 10, 8, 10, 12, 10, 13, 60])
    for t in r.trades:
        ws.append([t.day, "KOOP" if t.side == "BUY" else "VERKOOP", t.ticker, t.qty, t.price,
                   t.value_eur, t.costs_eur, t.pnl_eur, t.reason])

    tmp = path.with_suffix(".tmp.xlsx")
    wb.save(tmp)
    tmp.replace(path)


def write_chart(r: BacktestResult, path: Path) -> bool:
    try:
        import matplotlib
        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
    except ImportError:
        return False
    fig, ax = plt.subplots(figsize=(10, 5), dpi=120)
    styles = {"AGENT": ("#1F3A5F", 2.2), "IWDA.AS": ("#8C8C8C", 1.3), "SPY": ("#C9A227", 1.3)}
    for k, v in _all_series(r).items():
        color, width = styles.get(k, ("#555555", 1.2))
        ax.plot([d for d, _ in v], [x for _, x in v], label=NAMES.get(k, k), color=color, linewidth=width)
    for d, _ in r.halts:
        ax.axvline(d, color="#B03A2E", alpha=0.25, linewidth=1)
    ax.axhline(10_000, color="#000000", alpha=0.3, linewidth=0.8, linestyle="--")
    ax.set_title("Waarde van €10.000: trend-agent tegenover gewoon een ETF vasthouden")
    ax.set_ylabel("Waarde in euro")
    ax.grid(alpha=0.25)
    ax.legend(loc="upper left", frameon=False)
    for side in ("top", "right"):
        ax.spines[side].set_visible(False)
    fig.text(0.01, 0.01, "Rode lijnen: kill switch aangegaan. Nepgeld; geen voorspelling.", fontsize=8, alpha=0.7)
    fig.tight_layout()
    fig.savefig(path)
    plt.close(fig)
    return True


def write_all(r: BacktestResult, folder: Path, years: int) -> None:
    folder.mkdir(parents=True, exist_ok=True)
    write_chart(r, folder / "waardeverloop.png")
    write_excel(r, folder / "backtest.xlsx")
    write_markdown(r, folder / "RAPPORT.md", years)
