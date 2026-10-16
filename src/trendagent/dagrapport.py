"""Eén Excel-bestand met alles van de dag: rapporten/dagrapport.xlsx.

Tabbladen:
  Overzicht            stand van beide bots en wat ze uit de analyse halen
  Analyse              wat elke bot zag (markt, breedte, zeef, momentum) en zijn conclusies
  Verlies en risico    resultaat per positie, afstand tot de stop, remmen, en winst/verlies per belegging
  Kansen               de sterkste trends met uitleg, instapmoment en het plan van de bot
  Klaar voor echt geld de eisen en de stand per bot
  Waarde per dag       beide bots en de wereld-ETF
  Trades               alle aan- en verkopen van beide bots
  Controles            koersbronnen en waarschuwingen
  Grafieken            de candle-grafieken van vandaag
  Uitleg               wat de begrippen betekenen

Wordt elke dag opnieuw opgebouwd uit de bestanden die de bots al schreven. Verandert niets aan de portefeuilles.
"""

from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from . import insights
from .reports import EUR, PCT, _sheet
from .universe import BY_TICKER

# Winst groen, verlies rood (Excel-kleur 10 = groen).
RES_EUR = '[Color10]+€ #,##0.00;[Red]-€ #,##0.00;€ 0.00'
RES_PCT = '[Color10]+0.00%;[Red]-0.00%;0.00%'
GREEN = PatternFill("solid", fgColor="D9F2D9")
RED = PatternFill("solid", fgColor="F8D7D7")
WRAP = Alignment(wrap_text=True, vertical="top")
BOLD = Font(bold=True)
STATE = {"ok": "✅ gehaald", "bezig": "⏳ nog bezig", "nee": "❌ niet gehaald"}
BOTS = (("Investeerder", "investeerder", "investeerder.json"), ("Trend-agent", "trend", "portfolio.json"))


def _load(path: Path) -> dict:
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return {}


def _plain(line: str) -> str:
    return line.removeprefix("- ").replace("**", "")


MAANDEN = ["januari", "februari", "maart", "april", "mei", "juni", "juli", "augustus", "september", "oktober",
           "november", "december"]
DAGEN = ["maandag", "dinsdag", "woensdag", "donderdag", "vrijdag", "zaterdag", "zondag"]


def _nl(x: float, digits: int = 2) -> str:
    return f"{x:,.{digits}f}".replace(",", "X").replace(".", ",").replace("X", ".")


def _nl_date(day: str) -> str:
    from datetime import date
    try:
        d = date.fromisoformat(day)
    except ValueError:
        return day
    return f"{DAGEN[d.weekday()]} {d.day} {MAANDEN[d.month - 1]} {d.year}"


def _short_date(day: str) -> str:
    from datetime import date
    try:
        d = date.fromisoformat(day)
    except ValueError:
        return day
    return f"{d.day} {MAANDEN[d.month - 1][:3]}"


def _height(text, width_chars: float, size: int = 11) -> float:
    """Rijhoogte zodat afgebroken tekst helemaal zichtbaar is (Excel doet dat niet altijd zelf)."""
    import math
    if not isinstance(text, str) or not text:
        return 18
    per_line = max(width_chars * 11 / size * 1.05, 10)
    lines = sum(max(1, math.ceil(len(part) / per_line)) for part in text.split("\n"))
    return max(18, lines * size * 1.45 + 4)


def _fit_rows(ws) -> None:
    """Rijhoogtes voor tabellen met afgebroken tekst, op basis van de kolombreedtes."""
    from openpyxl.utils import get_column_letter
    for row in ws.iter_rows(min_row=2):
        h = 18
        for c in row:
            if isinstance(c.value, str) and len(c.value) > 12:
                width = ws.column_dimensions[get_column_letter(c.column)].width or 10
                h = max(h, _height(c.value, width))
        ws.row_dimensions[row[0].row].height = h


def _wrap_rows(ws, first: int = 2) -> None:
    for row in ws.iter_rows(min_row=first):
        for c in row:
            if isinstance(c.value, str):
                c.alignment = WRAP
    _fit_rows(ws)


def build(state_dir: Path, reports_dir: Path, day: str) -> Workbook:
    from .summary import _history_stats
    info = insights.load(state_dir)
    books = {key: _load(state_dir / fname) for _, key, fname in BOTS}
    health = _load(state_dir / "health.json")
    wb = Workbook()
    wb.remove(wb.active)

    # --- Overzicht: als een leesbare pagina, niet als brede tabel
    ws = wb.create_sheet("Overzicht")
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 30
    ws.column_dimensions["C"].width = 90
    ws["B1"] = f"Dagrapport {_nl_date(day)}"
    ws["B1"].font = Font(size=20, bold=True)
    ws["B2"] = "Twee oefenbots met elk €10.000 nepgeld. Geen beleggingsadvies."
    ws["B2"].font = Font(size=11, italic=True, color="52514E")
    row = 4
    for name, key, _ in BOTS:
        hist = books[key].get("equity_history", [])
        st = _history_stats(hist)
        start = hist[0]["equity"] if hist else None
        since = round(st["waarde"] - start, 2) if st and start else None
        label = "–" if since is None else ("WINST" if since > 0.005 else "VERLIES" if since < -0.005 else "GELIJK")
        realised = round(sum(t.get("pnl_eur") or 0 for t in books[key].get("trades", []) if t.get("side") == "SELL"), 2)
        ws.cell(row=row, column=2, value=name).font = Font(size=15, bold=True)
        row += 1
        lines = [("Waarde nu", st.get("waarde"), EUR, None)]
        if since is not None:
            lines.append((f"{label} sinds de start", since, RES_EUR, label))
            lines.append(("Sinds de start in %", st.get("sinds_start"), RES_PCT, None))
        if "vandaag" in st:
            lines.append(("Vandaag", st["vandaag"], RES_EUR, None))
        lines.append(("Al vastgezet door verkopen", realised, RES_EUR, None))
        if "vs_etf" in st:
            lines.append(("Beter of slechter dan de wereld-ETF", st["vs_etf"], RES_PCT, None))
        ready = info.get(key, {}).get("klaar", {}).get("oordeel")
        if ready:
            lines.append(("Klaar voor echt geld?", ready, None, None))
        for lab, val, fmt, flag in lines:
            ws.cell(row=row, column=2, value=lab).font = Font(size=12, color="52514E")
            c = ws.cell(row=row, column=3, value=val)
            c.font = Font(size=12, bold=flag is not None or lab == "Waarde nu")
            c.alignment = Alignment(horizontal="left", wrap_text=True)
            if fmt:
                c.number_format = fmt
            if flag in ("WINST", "VERLIES"):
                ws.cell(row=row, column=2).fill = GREEN if flag == "WINST" else RED
                ws.cell(row=row, column=2).font = Font(size=12, bold=True)
            row += 1
        row += 1

    def section(title: str, items: list[str]) -> None:
        nonlocal row
        ws.cell(row=row, column=2, value=title).font = Font(size=14, bold=True)
        row += 1
        for text in items:
            c = ws.cell(row=row, column=3, value=text)
            c.font = Font(size=12)
            c.alignment = WRAP
            ws.cell(row=row, column=2, value="•").alignment = Alignment(horizontal="right", vertical="top")
            ws.row_dimensions[row].height = _height(text, 90, 12)
            row += 1
        row += 1

    for name, key, _ in BOTS:
        d = info.get(key, {})
        if d.get("conclusies"):
            section(f"{name}: wat de bot uit de analyse haalt", d["conclusies"])
    done = [(name, t) for name, key, _ in BOTS for t in books[key].get("trades", []) if t.get("day") == day]
    section("Vandaag gekocht en verkocht", [
        f"{name} {'kocht' if t['side'] == 'BUY' else 'verkocht'} {t['ticker']} voor €{_nl(t['value_eur'])}"
        f" ({t['reason'][:90]})" for name, t in done] or ["Niets gekocht of verkocht."])
    ws.cell(row=row, column=2, value="Meer weten? Zie de tabbladen hieronder: Grafieken, Analyse, Verlies en risico, "
                                      "Kansen, Klaar voor echt geld.").font = Font(size=11, italic=True, color="52514E")
    ws.cell(row=row + 2, column=2).value = None
    ws._chart_row = row + 2   # hier komt de waardegrafiek

    # --- Analyse
    ws = _sheet(wb, "Analyse", ["Bot", "Onderdeel", "Wat de bot zag of concludeert"], [16, 22, 110])
    for name, key, _ in BOTS:
        d = info.get(key, {})
        for line in d.get("zag", []):
            text = _plain(line)
            head, _, rest = text.partition(":")
            ws.append([name, head.strip(), rest.strip()])
        for c in d.get("conclusies", []):
            ws.append([name, "Conclusie", c])
    _wrap_rows(ws)

    # --- Verlies en risico
    ws = _sheet(wb, "Verlies en risico", ["Bot", "Ticker", "Instap", "Koers nu", "Stop", "Resultaat €", "Resultaat %",
                                          "Afstand tot stop", "Als de stop geraakt wordt €", "In gewone taal"],
                [14, 12, 11, 11, 11, 13, 12, 15, 24, 70])
    tr = info.get("trend", {})
    for r in tr.get("posities", []):
        st = r["bij_stop_eur"]
        plain = (f"Staat €{_nl(abs(r['resultaat_eur']))} in de {'plus' if r['resultaat_eur'] >= 0 else 'min'}. "
                 f"Zakt de koers {_nl(abs(r['afstand_stop_pct'] or 0) * 100, 1)}% naar {_nl(r['stop'])}, dan verkoopt de bot "
                 + (f"met ongeveer €{_nl(-st)} verlies." if st < 0 else f"en is €{_nl(st)} winst vastgezet."))
        ws.append(["Trend-agent", r["ticker"], r["instap"], r["koers"], r["stop"], r["resultaat_eur"],
                   r["resultaat_pct"], r["afstand_stop_pct"], r["bij_stop_eur"], plain])
        row = ws.max_row
        for col, fmt in ((6, RES_EUR), (7, RES_PCT), (8, PCT), (9, RES_EUR)):
            ws.cell(row=row, column=col).number_format = fmt
    b = tr.get("remmen")
    if b:
        ws.append([])
        ws.append(["Trend-agent remmen", "Nu", "Grens"])
        ws.cell(row=ws.max_row, column=1).font = BOLD
        for label, now, limit in (("Daling vanaf top", b["daling_vanaf_top"], -b["grens_daling"]),
                                  ("Vandaag", b["vandaag"], -b["grens_dag"])):
            ws.append([label, now, limit])
            ws.cell(row=ws.max_row, column=2).number_format = PCT
            ws.cell(row=ws.max_row, column=3).number_format = PCT
        ws.append(["Verliestrades op rij", b["verlies_op_rij"], b["grens_op_rij"]])
    inv = info.get("investeerder", {})
    if inv.get("posities"):
        ws.append([])
        ws.append(["Investeerder", "Instrument", "Waarde", "Ingelegd", "Resultaat €", "Resultaat %"])
        for c in range(1, 7):
            ws.cell(row=ws.max_row, column=c).font = BOLD
        for r in inv["posities"]:
            ws.append(["", r["ticker"], r["waarde_eur"], r["kosten_eur"], r["resultaat_eur"], r["resultaat_pct"]])
            row = ws.max_row
            for col, fmt in ((3, EUR), (4, EUR), (5, RES_EUR), (6, RES_PCT)):
                ws.cell(row=row, column=col).number_format = fmt

    _wrap_rows(ws)

    # --- Kansen: per kans een blok dat je van boven naar beneden leest
    ws = wb.create_sheet("Kansen")
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 24
    ws.column_dimensions["C"].width = 95
    ws["B1"] = "Kansen: wat de trend-agent vandaag het sterkst vond"
    ws["B1"].font = Font(size=18, bold=True)
    ws["B2"] = ("Wat de oefenbot met nepgeld zou doen volgens zijn regels. Geen advies voor echt geld. "
                "Het instapmoment is een ongeteste vuistregel en wordt door de bot niet gebruikt.")
    ws["B2"].font = Font(size=11, italic=True, color="52514E")
    row = 4
    moment_fill = {"gunstig": GREEN, "opgelopen": RED}
    for i, k in enumerate(tr.get("kansen", []), 1):
        p = k["plan"]
        cur = p["valuta"]
        head = ws.cell(row=row, column=2, value=f"{i}. {k['naam']} ({k['ticker']})")
        head.font = Font(size=14, bold=True)
        row += 1
        if p["aantal"] >= 1:
            plan = (f"Koop {p['aantal']:g} {'stuk' if p['aantal'] == 1 else 'stuks'} rond {_nl(p['instap'])} {cur} "
                    f"(inzet €{_nl(p['inzet_eur'])}). Stop-loss op {_nl(p['stop'])} {cur} "
                    f"({_nl(p['stop_pct'] * 100, 1)}%). Gaat het mis, dan is het verlies ongeveer €{_nl(p['max_verlies_eur'])}.")
        else:
            plan = (f"Instap rond {_nl(p['instap'])} {cur}, stop-loss {_nl(p['stop'])} {cur}. "
                    "Er past geen heel stuk in het risicobudget van deze kleine oefenrekening.")
        for lab, text, fill in (("Instapmoment", f"{k['moment']}: {k['moment_uitleg']}", moment_fill.get(k["moment"])),
                                ("Waarom sterk", k["uitleg"], None),
                                ("Plan van de bot", plan, None),
                                ("Wat deed de bot?", k["status"], None)):
            lc = ws.cell(row=row, column=2, value=lab)
            lc.font = Font(size=12, color="52514E")
            lc.alignment = Alignment(vertical="top")
            if fill:
                lc.fill = fill
            c = ws.cell(row=row, column=3, value=text[0].upper() + text[1:] if text else text)
            c.font = Font(size=12, bold=lab == "Plan van de bot")
            c.alignment = WRAP
            ws.row_dimensions[row].height = _height(text, 95, 12)
            row += 1
        row += 1

    # --- Klaar voor echt geld
    ws = _sheet(wb, "Klaar voor echt geld", ["Bot", "Eis", "Stand", "Nu"], [16, 36, 16, 50])
    for name, key, _ in BOTS:
        k = info.get(key, {}).get("klaar")
        if not k:
            continue
        for label, state, detail in k["eisen"]:
            ws.append([name, label, STATE.get(state, state), detail])
        ws.append([name, "Oordeel", "", k["oordeel"]])
        ws.cell(row=ws.max_row, column=4).font = BOLD
        ws.append([])
    ws.append(["", "Eerst bewijzen, dan een oefenaccount bij een echte broker, dan een klein bedrag met jouw "
                   "goedkeuring per order. Die beslissing neem jij, niet de bot."])
    _wrap_rows(ws)

    # --- Waarde per dag
    ws = _sheet(wb, "Waarde per dag", ["Datum", "Investeerder", "Investeerder t.o.v. start", "Trend-agent",
                                       "Trend-agent t.o.v. start", "Wereld-ETF (koers in €)"], [14, 15, 22, 15, 22, 22])
    days: dict[str, list] = {}
    starts = []
    for col, (_, key, _) in enumerate(BOTS):
        hist = books[key].get("equity_history", [])
        starts.append(hist[0]["equity"] if hist else None)
        for e in hist:
            row = days.setdefault(e["day"], [None, None, None])
            row[col] = e["equity"]
            if "bench" in e:
                row[2] = e["bench"]
    order = sorted(days, key=lambda d: (d != "start", d))
    for d in order:
        inv_v, tr_v, bench = days[d]
        diff = [round(v - s0, 2) if v is not None and s0 else None for v, s0 in ((inv_v, starts[0]), (tr_v, starts[1]))]
        ws.append([d, inv_v, diff[0], tr_v, diff[1], bench])
        r = ws.max_row
        for col, fmt in ((2, EUR), (3, RES_EUR), (4, EUR), (5, RES_EUR)):
            ws.cell(row=r, column=col).number_format = fmt

    # --- Trades
    ws = _sheet(wb, "Trades", ["Datum", "Bot", "Ticker", "Koop/verkoop", "Aantal", "Prijs", "Waarde €", "Kosten €",
                               "Resultaat €", "Reden"], [12, 14, 10, 13, 10, 11, 12, 10, 12, 70])
    rows = [(t["day"], name, t) for name, key, _ in BOTS for t in books[key].get("trades", [])]
    for d, name, t in sorted(rows, key=lambda x: x[0]):
        ws.append([d, name, t["ticker"], "koop" if t["side"] == "BUY" else "verkoop", t["qty"], t["price"],
                   t["value_eur"], t["costs_eur"], t.get("pnl_eur"), t["reason"]])
        for col in (7, 8):
            ws.cell(row=ws.max_row, column=col).number_format = EUR
        ws.cell(row=ws.max_row, column=9).number_format = RES_EUR
    _wrap_rows(ws)

    # --- Controles
    ws = _sheet(wb, "Controles", ["Bot", "Controle"], [16, 120])
    from .summary import _plain_checks
    for name, key, _ in BOTS:
        for line in _plain_checks(name, health.get(key, {})):
            ws.append([name, line.replace("**", "")])
    _wrap_rows(ws)

    # --- Patronen in de grafieken
    try:
        from . import patronen
        folder = patronen.latest_folder(reports_dir)
        prow = patronen.per_ticker(folder) if folder else []
    except Exception:  # noqa: BLE001
        prow = []
    if prow:
        ws = wb.create_sheet("Patronen", 1)
        ws.sheet_view.showGridLines = False
        ws.column_dimensions["A"].width = 3
        for col, w in zip("BCDEFGH", (26, 13, 10, 10, 11, 60, 70)):
            ws.column_dimensions[col].width = w
        ws["B1"] = "Wat de grafieken laten zien"
        ws["B1"].font = Font(size=18, bold=True)
        row = 3
        for t in patronen.samenvatting(prow):
            c = ws.cell(row=row, column=2, value="• " + t)
            c.font = Font(size=12)
            ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=8)
            c.alignment = WRAP
            ws.row_dimensions[row].height = _height(t, 180, 12)
            row += 1
        row += 1
        heads = ["Aandeel", "Trend", "Week", "Maand", "3 maanden", "Opvallend", "Samenvatting"]
        for j, h in enumerate(heads, start=2):
            c = ws.cell(row=row, column=j, value=h)
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="1F3A5F")
        row += 1
        trend_fill = {"stijgend": GREEN, "dalend": RED}
        for r in prow:
            vals = [f"{r['naam']} ({r['ticker']})", r["trend"], r["week"], r["maand"], r["drie_mnd"],
                    r["patronen"] or "–", r["samenvatting"]]
            for j, v in enumerate(vals, start=2):
                c = ws.cell(row=row, column=j, value=v)
                c.alignment = WRAP
                if j in (4, 5, 6):
                    c.number_format = RES_PCT
            if r["trend"] in trend_fill:
                ws.cell(row=row, column=3).fill = trend_fill[r["trend"]]
            ws.row_dimensions[row].height = max(_height(r["samenvatting"], 70), _height(r["patronen"], 60))
            row += 1
        ws.cell(row=row + 1, column=2, value="Dit beschrijft wat er gebeurd is; het voorspelt niet wat de koersen "
                                             "morgen doen. De bots gebruiken het niet om te beslissen.").font = \
            Font(size=10, italic=True, color="52514E")

    # --- Uitleg
    ws = _sheet(wb, "Uitleg", ["Begrip", "Betekenis"], [26, 110])
    for row in (
        ("Wereld-ETF", "Een fonds met ongeveer 1.500 grote bedrijven wereldwijd (IWDA). Dit is de maatstaf: doet de bot het beter dan dit simpelweg vasthouden?"),
        ("200-daags gemiddelde", "De gemiddelde slotkoers van de laatste 200 handelsdagen. Erboven = stijgende trend op lange termijn."),
        ("Breedte", "Hoeveel van de gevolgde aandelen boven hun 200-daags gemiddelde staan. Hoog = de stijging wordt breed gedragen."),
        ("Beweeglijkheid", "Hoe hard de koersen op en neer gaan, vergeleken met het afgelopen jaar."),
        ("ATR", "De gemiddelde dagbeweging van een aandeel. De stop-loss ligt 2x de ATR onder de instap."),
        ("RSI", "Maat voor hoe hard een koers recent is gestegen (0-100). Boven 70-75 = mogelijk te snel opgelopen."),
        ("MACD", "Laat zien of een stijging versnelt (positief) of vertraagt (negatief)."),
        ("Stop-loss", "De prijs waarop de bot automatisch verkoopt om het verlies te beperken. Schuift mee omhoog als het goed gaat."),
        ("R", "Resultaat van een trade gedeeld door het risico bij instap. +2 R = twee keer zoveel gewonnen als je riskeerde."),
        ("Instapmoment", "Ongeteste vuistregel: gunstig = sterk maar niet ver weggelopen; opgelopen = ver boven het 20-daags gemiddelde."),
        ("Momentum", "Stijging over 12 maanden zonder de laatste maand. De investeerder kiest de 5 hoogste."),
        ("Kill switch", "Noodstop. Bij te veel verlies koopt de bot niets meer tot jij hem weer aanzet."),
    ):
        ws.append(list(row))
    _wrap_rows(ws)
    return wb


def _print_fit(ws) -> None:
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4


def _chart_title(png: Path) -> str:
    bot, _, ticker = png.stem.partition("-")
    inst = BY_TICKER.get(ticker)
    who = {"trend": "Trend-agent", "investeerder": "Investeerder"}.get(bot, bot)
    return f"{who}: {inst.name if inst else ticker} ({ticker})"


def write(state_dir: Path, reports_dir: Path, day: str) -> Path:
    reports_dir.mkdir(parents=True, exist_ok=True)
    path = reports_dir / "dagrapport.xlsx"
    tmp = path.with_suffix(".tmp.xlsx")
    wb = build(state_dir, reports_dir, day)
    for ws in wb.worksheets:   # afdrukken/PDF: altijd één pagina breed
        _print_fit(ws)
    wb.save(tmp)
    tmp.replace(path)
    try:   # echte Excel-grafieken: eerst de waarde van de bots, dan elke positie (trend-agent eerst)
        import json

        from . import charts, excel_charts
        books = {key: _load(state_dir / fname) for _, key, fname in BOTS}
        series = {name: [(e["day"], e["equity"]) for e in books[key].get("equity_history", []) if e["day"] != "start"]
                  for name, key, _ in BOTS}
        benches = [[(e["day"], e["bench"]) for e in books[key].get("equity_history", []) if "bench" in e]
                   for _, key, _ in BOTS]
        days = {d for pts in series.values() for d, _ in pts}
        labels = {d: _short_date(d) for d in days}
        folder = charts.chart_dir(reports_dir, day)
        files = sorted(folder.glob("*.json"), key=lambda p: (not p.name.startswith("trend"), p.name)) \
            if folder.exists() else []
        items = []
        for f in files:
            try:
                items.append(dict(json.loads(f.read_text(encoding="utf-8")), bot=f.stem.split("-", 1)[0]))
            except ValueError:
                continue
        excel_charts.add(path, items, {"series": series, "bench": max(benches, key=len), "labels": labels},
                         position=1)
    except Exception as exc:  # noqa: BLE001
        print(f"Excel-grafieken overgeslagen: {type(exc).__name__}: {exc}")
    return path
