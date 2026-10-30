"""Tests voor de trend-agent. Alles draait op verzonnen koersen, zonder internet."""

import json
import os
import tempfile
import unittest
from dataclasses import replace
from datetime import date, timedelta
from pathlib import Path
from unittest import mock

from trendagent.agent import TrendAgent
from trendagent.core.config import LiveTradingNotAllowedError, Settings, load_settings
from trendagent.core.kill_switch import KillSwitch
from trendagent.data import Bar, DataError, SimulatedProvider, validate_bars
from trendagent.portfolio import Portfolio
from trendagent.signals import analyse, atr, sma
from trendagent.universe import BY_TICKER, select

TODAY = date(2026, 10, 9)  # een vrijdag
UP = {"ASML.AS": 0.004, "NVDA": 0.004, "SAP.DE": 0.003, "SPY": 0.002}
FEW = ("ASML.AS", "NVDA", "SAP.DE", "SPY", "TTE.PA")


def settings(**kw) -> Settings:
    tmp = Path(tempfile.mkdtemp())
    base = Settings(state_dir=tmp / "state", reports_dir=tmp / "rapporten", tickers_override=FEW)
    return replace(base, **kw)


def flat_bars(n=260, price=100.0, end=TODAY):
    out, d = [], end - timedelta(days=n)
    for _ in range(n):
        d += timedelta(days=1)
        out.append(Bar(d, price, price * 1.01, price * 0.99, price, 1e6))
    return out


class ConfigTests(unittest.TestCase):
    def test_live_refused(self):
        with mock.patch.dict(os.environ, {"PAPER_TRADING": "false"}, clear=True):
            with self.assertRaises(LiveTradingNotAllowedError):
                load_settings()

    def test_risk_limits_capped(self):
        with mock.patch.dict(os.environ, {"RISK_PER_TRADE": "0.05"}, clear=True):
            with self.assertRaises(ValueError):
                load_settings()

    def test_unknown_ticker_refused(self):
        with self.assertRaises(ValueError):
            select(("GME",))


class IndicatorTests(unittest.TestCase):
    def test_sma(self):
        self.assertEqual(sma([1, 2, 3, 4], 2), 3.5)

    def test_atr_flat(self):
        self.assertAlmostEqual(atr(flat_bars()), 2.0, places=6)

    def test_flat_market_is_not_trending(self):
        a = analyse("X", flat_bars(), TODAY, settings())
        self.assertFalse(a.is_candidate)

    def test_stale_data_not_candidate(self):
        bars = SimulatedProvider(TODAY - timedelta(days=10), drifts=UP).history(BY_TICKER["NVDA"], 260)
        a = analyse("NVDA", bars, TODAY, settings())
        self.assertFalse(a.is_candidate)
        self.assertTrue(any("te oud" in r for r in a.reasons))

    def test_uptrend_is_candidate(self):
        bars = SimulatedProvider(TODAY, drifts=UP).history(BY_TICKER["NVDA"], 260)
        self.assertTrue(analyse("NVDA", bars, TODAY, settings()).is_candidate)

    def test_invalid_bars_rejected(self):
        bars = flat_bars()
        bars[5] = Bar(bars[5].day, 100, 90, 110, 100, 1e6)   # high < low
        with self.assertRaises(DataError):
            validate_bars(bars, "X")


class AgentTests(unittest.TestCase):
    def test_buys_trending_within_limits(self):
        s = settings()
        rep = TrendAgent(s, SimulatedProvider(TODAY, drifts=UP)).run(TODAY)
        pf = Portfolio.load(s.state_dir / "portfolio.json", s.starting_cash_eur)
        self.assertGreater(len(pf.positions), 0)
        self.assertLessEqual(len(pf.positions), s.max_positions)
        fx = {"EUR": 1.0, "USD": 0.92}
        self.assertLessEqual(pf.open_risk_eur(fx), s.starting_cash_eur * s.max_open_risk + 1e-6)
        for p in pf.positions.values():
            self.assertLessEqual(p.risk_eur, s.starting_cash_eur * s.risk_per_trade + 5)  # + kosten
            self.assertLess(p.stop, p.entry_price)
        self.assertTrue((s.state_dir / "report.md").exists())
        self.assertEqual(rep.status, "OK")

    def test_spreidt_over_sectoren(self):
        """Microsoft en SAP zijn allebei software: met hooguit 1 per sector koopt hij er maar één van."""
        from trendagent.universe import sector
        tk = ("MSFT", "SAP.DE", "NVDA", "SPY")
        drift = {"MSFT": 0.004, "SAP.DE": 0.0035, "NVDA": 0.003, "SPY": 0.002}
        uit = {}
        for grens in (1, 0):
            s = settings(max_positions=4, max_open_risk=0.02, max_per_sector=grens, starting_cash_eur=200_000,
                         tickers_override=tk)
            rep = TrendAgent(s, SimulatedProvider(TODAY, drifts=drift)).run(TODAY)
            pf = Portfolio.load(s.state_dir / "portfolio.json", s.starting_cash_eur)
            uit[grens] = ([sector(t) for t in pf.positions], rep.actions)
        secs, acties = uit[1]
        self.assertEqual(len(secs), len(set(secs)), acties)
        self.assertTrue(any("spreiden" in a for a in acties), acties)
        self.assertEqual(uit[0][0].count("Software en internet"), 2, uit[0][1])

    def test_never_buys_same_ticker_twice(self):
        s = settings()
        agent = TrendAgent(s, SimulatedProvider(TODAY, drifts=UP))
        agent.run(TODAY)
        agent.run(TODAY)
        pf = Portfolio.load(s.state_dir / "portfolio.json", s.starting_cash_eur)
        buys = [t.ticker for t in pf.trades if t.side == "BUY"]
        self.assertEqual(len(buys), len(set(buys)))

    def test_kill_switch_blocks_everything(self):
        s = settings()
        KillSwitch(s.state_dir / "kill_switch.json").activate("test", by="human")
        rep = TrendAgent(s, SimulatedProvider(TODAY, drifts=UP)).run(TODAY)
        pf = Portfolio.load(s.state_dir / "portfolio.json", s.starting_cash_eur)
        self.assertEqual(rep.status, "GESTOPT")
        self.assertEqual(pf.trades, [])

    def test_missing_fx_means_no_trade(self):
        s = settings()

        class NoFx(SimulatedProvider):
            def usd_to_eur(self):
                raise DataError("geen wisselkoers")

        rep = TrendAgent(s, NoFx(TODAY, drifts=UP)).run(TODAY)
        self.assertEqual(rep.status, "GEEN HANDEL")
        self.assertEqual(Portfolio.load(s.state_dir / "portfolio.json", 1).trades, [])

    def test_broken_data_stream_blocks_buys(self):
        s = settings()
        provider = SimulatedProvider(TODAY, drifts=UP, missing={"ASML.AS", "NVDA"})
        rep = TrendAgent(s, provider).run(TODAY)
        self.assertFalse(rep.buys_allowed)
        self.assertEqual(Portfolio.load(s.state_dir / "portfolio.json", 1).trades, [])

    def test_stop_loss_sells(self):
        s = settings()
        TrendAgent(s, SimulatedProvider(TODAY, drifts=UP)).run(TODAY)
        path = s.state_dir / "portfolio.json"
        pf = Portfolio.load(path, s.starting_cash_eur)
        ticker = next(iter(pf.positions))
        stop = pf.positions[ticker].stop

        class Crash(SimulatedProvider):
            def history(self, inst, days):
                bars = super().history(inst, days)
                if inst.ticker == ticker:
                    last = bars[-1]
                    crash = Bar(last.day, stop * 0.98, stop * 0.99, stop * 0.95, stop * 0.96, last.volume)
                    bars = bars[:-1] + [crash]
                return bars

        nxt = TODAY + timedelta(days=3)  # maandag
        TrendAgent(s, Crash(nxt, drifts=UP)).run(nxt)
        pf = Portfolio.load(path, s.starting_cash_eur)
        sells = [t for t in pf.trades if t.side == "SELL" and t.ticker == ticker]
        self.assertEqual(len(sells), 1)
        self.assertIn("stop-loss", sells[0].reason)
        self.assertLess(sells[0].pnl_eur, 0)

    def test_drawdown_triggers_kill_switch(self):
        s = settings()
        path = s.state_dir / "portfolio.json"
        pf = Portfolio(cash_eur=9_000.0, equity_history=[{"day": "start", "equity": 10_000.0}])
        pf.save(path)
        rep = TrendAgent(s, SimulatedProvider(TODAY, drifts=UP)).run(TODAY)
        self.assertEqual(rep.status, "GESTOPT")
        self.assertTrue(KillSwitch(s.state_dir / "kill_switch.json").is_active())

    def test_three_losses_trigger_kill_switch(self):
        s = settings()
        Portfolio(cash_eur=10_000.0, consecutive_losses=3).save(s.state_dir / "portfolio.json")
        rep = TrendAgent(s, SimulatedProvider(TODAY, drifts=UP)).run(TODAY)
        self.assertEqual(rep.status, "GESTOPT")

    def test_audit_log_intact(self):
        s = settings()
        agent = TrendAgent(s, SimulatedProvider(TODAY, drifts=UP))
        agent.run(TODAY)
        self.assertTrue(agent.audit.verify())
        events = [json.loads(l)["event"] for l in (s.state_dir / "audit.jsonl").read_text().splitlines()]
        self.assertIn("paper_trade", events)

    def test_reports_written(self):
        from openpyxl import load_workbook
        s = settings()
        TrendAgent(s, SimulatedProvider(TODAY, drifts=UP)).run(TODAY)
        self.assertTrue((s.reports_dir / f"{TODAY.isoformat()}.md").exists())
        wb = load_workbook(s.reports_dir / "trend-agent.xlsx")
        self.assertEqual(wb.sheetnames, ["Overzicht", "Posities", "Trades", "Trends", "Uitleg"])
        self.assertGreater(wb["Trades"].max_row, 1)


if __name__ == "__main__":
    unittest.main()


class AutoHervatTests(unittest.TestCase):
    def test_rem_van_risk_manager_gaat_na_n_dagen_vanzelf_uit(self):
        s = settings(auto_hervat_dagen=2)
        prov = SimulatedProvider(TODAY, drifts=UP)
        agent = TrendAgent(s, prov)
        ks = KillSwitch(s.state_dir / "kill_switch.json")
        ks.activate("3 verliestrades op rij", by="risk_manager")
        agent.run(TODAY - timedelta(days=1))
        self.assertTrue(ks.is_active())                 # dag 1: nog aan
        rep = agent.run(TODAY)
        self.assertFalse(ks.is_active())                # dag 2: vanzelf weer uit
        self.assertTrue(any("vanzelf weer uit" in w for w in rep.warnings))

    def test_rem_van_mens_blijft_aan(self):
        s = settings(auto_hervat_dagen=1)
        agent = TrendAgent(s, SimulatedProvider(TODAY, drifts=UP))
        ks = KillSwitch(s.state_dir / "kill_switch.json")
        ks.activate("handmatig", by="human")
        agent.run(TODAY)
        agent.run(TODAY + timedelta(days=3))
        self.assertTrue(ks.is_active())

    def test_standaard_uit(self):
        s = settings()
        agent = TrendAgent(s, SimulatedProvider(TODAY, drifts=UP))
        ks = KillSwitch(s.state_dir / "kill_switch.json")
        ks.activate("3 verliestrades op rij", by="risk_manager")
        for i in range(4):
            agent.run(TODAY - timedelta(days=3 - i))
        self.assertTrue(ks.is_active())


class RestEtfTests(unittest.TestCase):
    def test_ongebruikt_geld_gaat_in_de_etf_en_komt_vrij_voor_trends(self):
        s = settings(rest_etf="SPY", max_positions=2)
        agent = TrendAgent(s, SimulatedProvider(TODAY, drifts=UP))
        rep = agent.run(TODAY)
        pf = Portfolio.load(s.state_dir / "portfolio.json", s.starting_cash_eur)
        self.assertGreater(pf.geparkeerd_qty, 0)
        self.assertGreaterEqual(pf.cash_eur, 0)
        self.assertLess(pf.cash_eur, 0.05 * rep.equity_eur + 700)      # hooguit de buffer plus één stuk
        # geparkeerd geld is geen positie en geen trade
        self.assertLessEqual(len(pf.positions), 2)
        # meer plekken: de nieuwe aankoop wordt betaald door een deel van de ETF te verkopen
        voor = pf.geparkeerd_qty
        s2 = replace(s, max_positions=4)
        agent2 = TrendAgent(s2, SimulatedProvider(TODAY + timedelta(days=3), drifts=UP))
        agent2.run(TODAY + timedelta(days=3))
        pf2 = Portfolio.load(s.state_dir / "portfolio.json", s.starting_cash_eur)
        if len(pf2.positions) > len(pf.positions):
            self.assertLess(pf2.geparkeerd_qty, voor + 1e-9)
        self.assertGreaterEqual(pf2.cash_eur, 0)

    def test_remmen_kijken_niet_naar_de_etf(self):
        pf = Portfolio(cash_eur=1000.0)
        pf.parkeer("SPY", "USD", 10, 100.0, 1.0, 1.0)
        self.assertAlmostEqual(pf.rem_waarde(pf.equity({"SPY": 80.0}, {"USD": 1.0, "EUR": 1.0}),
                                             {"SPY": 80.0}, {"USD": 1.0, "EUR": 1.0}), 1000.0)
        pf.ontparkeer(5, 120.0, 1.0, 1.0)
        self.assertAlmostEqual(pf.geparkeerd_kosten, 500.5)
        self.assertAlmostEqual(pf.cash_eur, 1000 - 1001 + 599)

    def test_aankoop_wordt_betaald_uit_de_etf(self):
        s = settings(rest_etf="SPY", max_positions=4)
        s.state_dir.mkdir(parents=True)
        pf = Portfolio(cash_eur=10000.0)
        pf.parkeer("SPY", "USD", 30, 330.0, 0.9, 1.0)
        pf.save(s.state_dir / "portfolio.json")
        rep = TrendAgent(s, SimulatedProvider(TODAY, drifts=UP)).run(TODAY)
        pf = Portfolio.load(s.state_dir / "portfolio.json", s.starting_cash_eur)
        self.assertTrue(any("verkocht uit de rest-ETF" in a for a in rep.actions))
        self.assertLess(pf.geparkeerd_qty, 30)
        self.assertGreaterEqual(pf.cash_eur, 0)
        self.assertEqual(rep.status, "OK")
